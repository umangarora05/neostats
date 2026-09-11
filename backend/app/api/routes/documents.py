import io
import json

from docx import Document
from fastapi import APIRouter, UploadFile, File, Form, Depends, HTTPException
from fastapi.responses import Response
from pymongo.database import Database
from typing import List
from app.core.database import get_db
from app.services.document_service import process_document
from app.repositories import document_repository
from app.schemas.document import DocumentResponse

router = APIRouter()


def _document_response_data(doc: dict) -> dict:
    return DocumentResponse(
        document_name=doc.get("document_name"),
        document_type=doc.get("document_type"),
        processing_status=doc.get("processing_status"),
        file_validation=doc.get("file_validation"),
        extracted_data=doc.get("extracted_data") or {},
        validation=doc.get("validation") or {"checks": [], "overall_status": "UNKNOWN", "issues": []},
        processing_metadata=doc.get("processing_metadata")
    ).model_dump()


def _download_name(document_name: str, extension: str) -> str:
    base_name = document_name.rsplit(".", 1)[0]
    safe_name = "".join(char if char.isalnum() or char in "-_" else "_" for char in base_name)
    return f"{safe_name}.{extension}"


def _add_value_paragraph(document: Document, label: str, value) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(str(value))


def _create_word_document(data: dict) -> bytes:
    document = Document()
    document.add_heading(f"Document Result: {data['document_name']}", level=1)
    document.add_heading("Processing Summary", level=2)
    _add_value_paragraph(document, "Document Type", data["document_type"].replace("_", " ").title())
    _add_value_paragraph(document, "Processing Status", data["processing_status"])
    _add_value_paragraph(document, "Processed At", data["processing_metadata"]["processed_at"])
    _add_value_paragraph(document, "Processing Time", f"{data['processing_metadata']['processing_time_ms']} ms")

    document.add_heading("Extracted Data", level=2)
    document.add_paragraph(json.dumps(data["extracted_data"], indent=2, ensure_ascii=False))

    document.add_heading("Financial Validation", level=2)
    _add_value_paragraph(document, "Overall Status", data["validation"]["overall_status"])
    for issue in data["validation"].get("issues", []):
        document.add_paragraph(issue, style="List Bullet")

    document.add_heading("JSON Output", level=2)
    document.add_paragraph(json.dumps(data, indent=2, ensure_ascii=False))

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()

@router.post("/process", response_model=DocumentResponse)
async def process_document_endpoint(
    file: UploadFile = File(...),
    document_type: str = Form(...),
    db: Database = Depends(get_db)
):
    if document_type not in ["invoice", "balance_sheet", "profit_and_loss", "cash_flow_statement"]:
        raise HTTPException(status_code=400, detail="Invalid document_type")
        
    mime_type = file.content_type
    if mime_type not in ["application/pdf", "image/jpeg", "image/png", "image/jpg"]:
        from fastapi.responses import JSONResponse
        return JSONResponse(
            status_code=400,
            content={
                "error": {
                    "code": "UNSUPPORTED_FILE_TYPE",
                    "message": "Only PDF / JPG / PNG documents are supported."
                }
            }
        )

    file_content = await file.read()
    response_data = await process_document(db, file.filename, file_content, mime_type, document_type)
    
    return response_data


@router.get("/{document_name}/json")
def download_document_json(document_name: str, db: Database = Depends(get_db)):
    doc = document_repository.get_document_by_name(db, document_name)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    data = _document_response_data(doc)
    return Response(
        content=json.dumps(data, indent=2, ensure_ascii=False),
        media_type="application/json",
        headers={
            "Content-Disposition": f'attachment; filename="{_download_name(document_name, "json")}"'
        },
    )


@router.get("/{document_name}/word")
def download_document_word(document_name: str, db: Database = Depends(get_db)):
    doc = document_repository.get_document_by_name(db, document_name)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    data = _document_response_data(doc)
    return Response(
        content=_create_word_document(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{_download_name(document_name, "docx")}"'
        },
    )


@router.get("/{document_name}", response_model=DocumentResponse)
def get_document(document_name: str, db: Database = Depends(get_db)):
    doc = document_repository.get_document_by_name(db, document_name)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    
    return _document_response_data(doc)

@router.get("", response_model=List[DocumentResponse])
def list_documents(skip: int = 0, limit: int = 100, db: Database = Depends(get_db)):
    docs = document_repository.list_documents(db, skip=skip, limit=limit)
    return [
        _document_response_data(doc) for doc in docs
    ]
